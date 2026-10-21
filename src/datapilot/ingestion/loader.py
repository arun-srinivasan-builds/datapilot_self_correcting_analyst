"""CSV/XLSX loading with validation. Schema-agnostic: no column names are assumed."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from datapilot.ingestion import validate
from datapilot.ingestion.models import Dataset, IngestionError

DEFAULT_MAX_UPLOAD_MB = 25
DEFAULT_MAX_ROWS = 1_000_000
DEFAULT_MAX_COLUMNS = 200


def list_sheets(path: Path, max_upload_mb: int = DEFAULT_MAX_UPLOAD_MB) -> tuple[str, ...]:
    """Visible sheet names for XLSX; empty tuple for CSV."""
    kind = validate.check_file(path, max_upload_mb)
    if kind == "csv":
        return ()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(ws.title for ws in wb.worksheets if ws.sheet_state == "visible")
    finally:
        wb.close()


def _fingerprint(path: Path, sheet: str | None) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    digest.update((sheet or "").encode())
    return digest.hexdigest()


def load_dataset(
    path: Path,
    sheet: str | None = None,
    *,
    max_upload_mb: int = DEFAULT_MAX_UPLOAD_MB,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
) -> Dataset:
    """Load one CSV or one sheet of an XLSX workbook. Raises IngestionError on rejection."""
    path = Path(path)
    kind = validate.check_file(path, max_upload_mb)
    sheets: tuple[str, ...] = ()
    chosen: str | None = None
    try:
        if kind == "csv":
            encoding, delimiter = validate.sniff_csv(path)
            validate.check_header(validate.header_of_csv(path, encoding, delimiter))
            frame = pd.read_csv(
                path, encoding=encoding, sep=delimiter, nrows=max_rows + 1, low_memory=False
            )
        else:
            sheets = list_sheets(path, max_upload_mb)
            if not sheets:
                raise IngestionError("no_sheets", "The workbook has no visible sheets.")
            chosen = sheet or sheets[0]
            if chosen not in sheets:
                raise IngestionError("unknown_sheet", f"Sheet '{chosen}' was not found.")
            frame = pd.read_excel(path, sheet_name=chosen, engine="openpyxl")
            validate.check_header(_xlsx_header(path, chosen))
    except IngestionError:
        raise
    except (pd.errors.ParserError, pd.errors.EmptyDataError, ValueError, KeyError) as exc:
        raise IngestionError("unreadable", f"The file could not be parsed: {exc}") from exc

    if frame.empty or frame.shape[1] == 0:
        raise IngestionError("empty_data", "The file has a header but no data rows.")
    if len(frame) > max_rows:
        raise IngestionError("too_many_rows", f"More than {max_rows:,} rows are not supported.")
    if frame.shape[1] > max_columns:
        raise IngestionError(
            "too_many_columns", f"More than {max_columns} columns are not supported."
        )
    frame.columns = [str(c).strip() for c in frame.columns]
    return Dataset(
        frame=frame,
        name=path.name,
        sheet=chosen,
        sheets=sheets,
        fingerprint=_fingerprint(path, chosen),
        size_bytes=path.stat().st_size,
    )


def _xlsx_header(path: Path, sheet: str) -> list[object]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        row = next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [c for c in row if c is not None] or [None]
    finally:
        wb.close()
